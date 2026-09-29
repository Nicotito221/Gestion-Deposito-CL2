import argparse
import csv
import datetime
import os
import re
import tkinter as tk
import unicodedata
from dataclasses import dataclass, field
from tkinter import filedialog, messagebox
from typing import Any, Dict, List, Optional, Set, Tuple, Union

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SPANISH_MONTHS = {
    1: 'ene', 2: 'feb', 3: 'mar', 4: 'abr', 5: 'may', 6: 'jun',
    7: 'jul', 8: 'ago', 9: 'sept', 10: 'oct', 11: 'nov', 12: 'dic'
}

MONTHS_LOOKUP = {
    'ene': 1, 'enero': 1, 'jan': 1,
    'feb': 2, 'febrero': 2,
    'mar': 3, 'marzo': 3,
    'abr': 4, 'abril': 4, 'apr': 4,
    'may': 5, 'mayo': 5,
    'jun': 6, 'junio': 6,
    'jul': 7, 'julio': 7,
    'ago': 8, 'agosto': 8, 'aug': 8,
    'sep': 9, 'sept': 9, 'septiembre': 9,
    'oct': 10, 'octubre': 10,
    'nov': 11, 'noviembre': 11,
    'dic': 12, 'diciembre': 12, 'dec': 12
}

DEFAULT_BOM_CANDIDATES = [
    r'C:\Users\nangerossa\Desktop\Cargar orden suministro\Interfaz.xlsm',
    r'C:\Users\nangerossa\Desktop\Cargar orden suministro\interfax.xlsm',
    r'C:\Users\nangerossa\Desktop\Cargar orden suministro\interfax.xslm',
    r'C:\Users\nangerossa\Desktop\Cargar orden suministro\Interfaz.xlsx',
]

DEFAULT_TENTATIVE_DATES = {
    10: '19.10.2026',
    11: '19.11.2026',
    12: '14.12.2026',
    1: '18.01.2027',
    2: '15.02.2027'
}

DEFAULT_MATERIALS_CANDIDATES = [
    r'C:\Users\nangerossa\Desktop\PRD - Septiembre 2026 - IASS015922 (2).csv',
    r'C:\Users\nangerossa\Desktop\Cargar orden suministro\PRD - Septiembre 2026 - IASS015922 (2).csv',
    r'C:\Users\nangerossa\Desktop\PRD - Septiembre 2026.xlsx',
    r'C:\Users\nangerossa\Desktop\Cargar orden suministro\PRD - Septiembre 2026.xlsx',
]

def clean_code(raw: Any) -> str:
    if raw is None:
        return ''
    raw_str = str(raw).replace('º', '').replace('°', '')
    nfkd = unicodedata.normalize('NFKD', raw_str)
    no_accents = ''.join([c for c in nfkd if not unicodedata.combining(c)])
    return re.sub(r'[^A-Za-z0-9]', '', no_accents).upper()


def get_field(row: Dict[str, Any], *candidates: str, default: Any = '') -> Any:
    for c in candidates:
        if c in row and row[c] is not None and str(row[c]).strip() != '':
            return row[c]

    cleaned_candidates = [clean_code(c) for c in candidates]
    for k, v in row.items():
        if v is None or str(v).strip() == '':
            continue
        k_clean = clean_code(k)
        for c_clean in cleaned_candidates:
            if k_clean == c_clean:
                return v
    return default


def format_part_number(raw: Any) -> str:
    if not raw:
        return ''
    val = str(raw).strip()

    for suffix in ('-SJ', '-UJ', '-MTA'):
        if val.upper().endswith(suffix):
            val = val[:-len(suffix)]
            break

    if re.match(r'^[A-Z0-9]{5}-[A-Z0-9]{5}(-[A-Z0-9]{2})?$', val, re.IGNORECASE):
        return val.upper()

    clean = clean_code(val)
    for suffix in ('SJ', 'UJ', 'MTA'):
        if clean.endswith(suffix) and len(clean) > 12:
            clean = clean[:-len(suffix)]
            break

    if len(clean) == 10:
        return f'{clean[:5]}-{clean[5:]}'
    elif len(clean) == 12:
        return f'{clean[:5]}-{clean[5:10]}-{clean[10:]}'
    elif len(clean) > 12 and (clean.endswith('SJ') or clean.endswith('UJ')):
        base = clean[:-2]
        if len(base) == 10:
            return f'{base[:5]}-{base[5:]}'
        elif len(base) == 12:
            return f'{base[:5]}-{base[5:10]}-{base[10:]}'

    return val


def clean_numeric(val: Any) -> float:
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    val_str = str(val).strip()
    if not val_str or val_str.startswith('#') or val_str.upper() in ('ERROR', 'REF', 'N/A', 'NULL', 'NONE', '-'):
        return 0.0
    if val_str.endswith('-'):
        val_str = '-' + val_str[:-1]
    if ',' in val_str and '.' in val_str:
        val_str = val_str.replace('.', '').replace(',', '.')
    else:
        val_str = val_str.replace(',', '.')
    try:
        return float(val_str)
    except ValueError:
        return 0.0


def normalize_header_cell(val: Any) -> str:
    if val is None:
        return ''
    if isinstance(val, (datetime.datetime, datetime.date)):
        m_name = SPANISH_MONTHS.get(val.month, 'mes')
        return f'{val.day}-{m_name}'

    val_str = str(val).strip()
    m_dot = re.match(r'^(\d{1,2})\.(\d{1,2})\.(\d{2,4})', val_str)
    if m_dot:
        day = int(m_dot.group(1))
        month = int(m_dot.group(2))
        m_name = SPANISH_MONTHS.get(month, 'mes')
        return f'{day}-{m_name}'

    m_iso = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', val_str)
    if m_iso:
        day = int(m_iso.group(3))
        month = int(m_iso.group(2))
        m_name = SPANISH_MONTHS.get(month, 'mes')
        return f'{day}-{m_name}'

    m_slash = re.match(r'^(\d{1,2})[/.-](\d{1,2})([/.-]\d{2,4})?$', val_str)
    if m_slash:
        day = int(m_slash.group(1))
        month = int(m_slash.group(2))
        if 1 <= month <= 12 and 1 <= day <= 31:
            m_name = SPANISH_MONTHS.get(month, 'mes')
            return f'{day}-{m_name}'

    return val_str


def is_daily_column(header: str) -> bool:
    h = header.strip().lower()
    if not h:
        return False
    if 'total' in h or 'tentativo' in h or 'mes' in h:
        return False
    if re.match(r'^\d{1,2}-[a-z]+', h):
        return True
    if re.match(r'^\d{1,2}$', h):
        num = int(h)
        return 1 <= num <= 31
    return False


def is_future_tentative_column(header: str, firm_month: int = 9) -> bool:
    h_clean = header.strip().lower()
    if is_daily_column(header):
        return False
    for m_key, m_val in MONTHS_LOOKUP.items():
        if m_key in h_clean:
            if m_val == firm_month:
                return False
            return True
    return False


def read_table_rows(filepath: str, preferred_sheet_hint: str = '') -> Tuple[List[str], List[Dict[str, Any]], int]:
    ext = os.path.splitext(filepath)[1].lower()
    detected_year = datetime.datetime.now().year

    if ext in ('.xlsx', '.xlsm', '.xltx', '.xltm'):
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)

        target_ws = None
        if preferred_sheet_hint:
            for sname in wb.sheetnames:
                if preferred_sheet_hint.lower() in sname.lower():
                    target_ws = wb[sname]
                    break

        if target_ws is None:
            for sname in wb.sheetnames:
                ws = wb[sname]
                for row in ws.iter_rows(values_only=True, max_row=6):
                    if not row:
                        continue
                    str_cells = [clean_code(c) for c in row if c is not None]
                    if any('PARTE' in c or 'PART' in c or 'SUPPLIER' in c or 'TBAR' in c or 'HIJO' in c or 'SECTOR' in c for c in str_cells):
                        target_ws = ws
                        break
                if target_ws is not None:
                    break

        if target_ws is None:
            target_ws = wb.active

        rows_iter = target_ws.iter_rows(values_only=True)
        raw_headers: Optional[Tuple[Any, ...]] = None
        for row in rows_iter:
            if not row or not any(row):
                continue
            str_cells = [clean_code(c) for c in row if c is not None]
            if any('PARTE' in c or 'PART' in c or 'SUPPLIER' in c or 'TBAR' in c or 'HIJO' in c or 'SECTOR' in c for c in str_cells):
                raw_headers = row
                break

        if raw_headers is None:
            wb.close()
            return [], [], detected_year

        for h in raw_headers:
            if isinstance(h, (datetime.datetime, datetime.date)):
                detected_year = h.year
                break
            elif isinstance(h, str):
                m = re.search(r'202\d', h)
                if m:
                    detected_year = int(m.group(0))
                    break

        headers = [normalize_header_cell(h) for h in raw_headers]
        records: List[Dict[str, Any]] = []
        for row in rows_iter:
            if not any(row):
                continue
            rec = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
            records.append(rec)
        wb.close()
        return headers, records, detected_year

    elif ext in ('.csv', '.txt'):
        with open(filepath, mode='r', encoding='utf-8-sig') as f:
            sample = f.read(2048)
            f.seek(0)
            delimiter = '\t' if '\t' in sample else ','
            reader = csv.reader(f, delimiter=delimiter)
            raw_headers = None
            for row in reader:
                if not row or not any(row):
                    continue
                str_cells = [clean_code(c) for c in row if c]
                if any('PARTE' in c or 'PART' in c or 'SUPPLIER' in c or 'TBAR' in c or 'HIJO' in c or 'SECTOR' in c for c in str_cells):
                    raw_headers = row
                    break

            if not raw_headers:
                return [], [], detected_year

            for h in raw_headers:
                m = re.search(r'202\d', h)
                if m:
                    detected_year = int(m.group(0))
                    break

            headers = [normalize_header_cell(h) for h in raw_headers]
            records = []
            for row in reader:
                if not any(row):
                    continue
                rec = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
                records.append(rec)
            return headers, records, detected_year

    return [], [], detected_year


@dataclass
class ProductionItem:
    canonical_pn: str
    n_parte: str = ''
    n_parte_sep: str = ''
    codigo_tasa_tbar: str = ''
    pn_tbar: str = ''
    sd_tbar: str = ''
    descripcion: str = ''
    n_interno: str = ''
    tipo_caja: str = ''
    cant_x_caja: str = '1'
    dock: str = ''
    destino: str = 'PROD'
    daily_values: Dict[str, float] = field(default_factory=dict)
    tentative_values: Dict[str, float] = field(default_factory=dict)


@dataclass
class BomItem:
    sector: str = ''
    sd: str = ''
    codigo_tasa: str = ''
    codigo_sap_padre: str = ''
    hijo: str = ''
    uso: float = 1.0
    clave: str = ''
    daily_calculated: Dict[str, float] = field(default_factory=dict)


@dataclass
class MaterialMetadata:
    canonical_hijo: str
    n_parte: str = ''
    codigo_tasa_tbar: str = ''
    descripcion: str = ''
    n_interno: str = ''
    tipo_caja: str = ''
    cant_x_caja: str = '1'
    dock: str = ''
    destino: str = 'PROD'
    centro: str = '1000'


class ProductionConsolidator:
    def __init__(self) -> None:
        self.demand_items: Dict[str, ProductionItem] = {}
        self.demand_lookup: Dict[str, ProductionItem] = {}
        self.bom_items: List[BomItem] = []
        self.materials_metadata: Dict[str, MaterialMetadata] = {}
        self.daily_columns: List[str] = []
        self.tentative_columns_f1: List[str] = []
        self.tentative_columns_f2: List[str] = []
        self.custom_tentative_dates: Dict[int, str] = dict(DEFAULT_TENTATIVE_DATES)
        self.plan_year: int = datetime.datetime.now().year

    def register_demand_item(self, item: ProductionItem) -> None:
        self.demand_items[item.canonical_pn] = item
        for key in (item.canonical_pn, clean_code(item.n_parte), clean_code(item.n_parte_sep), clean_code(item.codigo_tasa_tbar), clean_code(item.pn_tbar)):
            if key:
                self.demand_lookup[key] = item

    def get_or_create_demand(self, canonical_pn: str) -> ProductionItem:
        if canonical_pn not in self.demand_items:
            self.demand_items[canonical_pn] = ProductionItem(canonical_pn=canonical_pn)
        return self.demand_items[canonical_pn]

    def load_mapping_file(self, filepath: str) -> None:
        if not os.path.exists(filepath):
            return
        _, records, year = read_table_rows(filepath, preferred_sheet_hint='map')
        if year:
            self.plan_year = year
        for row in records:
            raw_pn = get_field(row, 'Nº Parte', 'N° Parte', 'Part Number', 'PN TBAR', 'Material')
            canonical = clean_code(raw_pn)
            if not canonical:
                continue
            item = self.get_or_create_demand(canonical)
            item.n_parte = format_part_number(raw_pn)
            item.n_parte_sep = format_part_number(get_field(row, 'Nº Parte Sep', 'N° Parte Sep', default=item.n_parte_sep))
            item.codigo_tasa_tbar = format_part_number(get_field(row, 'Codigo TASA+TBAR', 'Codigo TASA', default=item.codigo_tasa_tbar))
            item.pn_tbar = format_part_number(get_field(row, 'PN TBAR', default=item.pn_tbar))
            item.sd_tbar = str(get_field(row, 'SD TBAR', 'SD', 'Sd', default=item.sd_tbar))
            self.register_demand_item(item)

    def load_bom_from_db(self) -> None:
        try:
            from app.core.config import get_db_connection
            conn = get_db_connection()
        except Exception:
            return

        cursor = None
        try:
            cursor = conn.cursor()
            cursor.execute(
                """
                SELECT 
                    s.sector,
                    s.sd,
                    s.codigo_tasa,
                    s.codigo_sap_padre,
                    s.id_material_hijo,
                    s.cantidad,
                    COALESCE(m.nombre, ''),
                    COALESCE(m.centro, '1000')
                FROM dbo.set_componente s
                LEFT JOIN dbo.material m ON s.id_material_hijo = m.id_material
                WHERE s.activo = 1
                """
            )
            rows = cursor.fetchall()
            for r in rows:
                sector = str(r[0] or 'Seat')
                sd = str(r[1] or '')
                codigo_tasa_raw = str(r[2] or '')
                padre_raw = str(r[3] or codigo_tasa_raw)
                hijo_raw = str(r[4] or '')
                uso_val = float(r[5] or 1.0)
                desc = str(r[6] or '')
                centro = str(r[7] or '1000')

                if not hijo_raw:
                    continue

                formatted_hijo = format_part_number(hijo_raw)
                formatted_padre = format_part_number(padre_raw)
                formatted_tasa = format_part_number(codigo_tasa_raw)

                b_item = BomItem(
                    sector=sector,
                    sd=sd,
                    codigo_tasa=formatted_tasa or codigo_tasa_raw,
                    codigo_sap_padre=formatted_padre or padre_raw,
                    hijo=formatted_hijo or hijo_raw,
                    uso=uso_val,
                    clave=f"{codigo_tasa_raw}{hijo_raw}"
                )
                self.bom_items.append(b_item)

                canon_hijo = clean_code(hijo_raw)
                if canon_hijo not in self.materials_metadata:
                    self.materials_metadata[canon_hijo] = MaterialMetadata(
                        canonical_hijo=canon_hijo,
                        n_parte=formatted_hijo,
                        codigo_tasa_tbar=formatted_hijo,
                        descripcion=desc,
                        centro=centro,
                    )
        except Exception:
            pass
        finally:
            if cursor:
                try:
                    cursor.close()
                except Exception:
                    pass
            try:
                conn.close()
            except Exception:
                pass

    def load_bom_file(self, filepath: str) -> None:
        if not os.path.exists(filepath):
            return
        _, records, year = read_table_rows(filepath, preferred_sheet_hint='bom')
        if year:
            self.plan_year = year
        for row in records:
            codigo_tasa_raw = str(get_field(row, 'Codigo TASA', 'Codigo TASA+TBAR', 'TASA', 'Codigo Padre', default=''))
            padre_raw = str(get_field(row, 'Codigo SAP (Padre)', 'Codigo SAP', 'Padre', default=codigo_tasa_raw))
            hijo_raw = str(get_field(row, 'Hijo', 'Nº Parte', 'N° Parte', 'Part Number', 'Componente', 'Material', default=''))
            uso_val = clean_numeric(get_field(row, 'Uso', 'Cantidad', 'Ctd.componente', default=1)) or 1.0

            if not hijo_raw:
                continue

            formatted_hijo = format_part_number(hijo_raw)
            formatted_padre = format_part_number(padre_raw)
            formatted_tasa = format_part_number(codigo_tasa_raw)

            b_item = BomItem(
                sector=str(get_field(row, 'Sector', default='Seat')),
                sd=str(get_field(row, 'Sd', 'SD', 'SD TBAR', default='')),
                codigo_tasa=formatted_tasa or codigo_tasa_raw,
                codigo_sap_padre=formatted_padre or padre_raw,
                hijo=formatted_hijo or hijo_raw,
                uso=uso_val,
                clave=str(get_field(row, 'Clave', default=f'{codigo_tasa_raw}{hijo_raw}'))
            )
            self.bom_items.append(b_item)

            canon_hijo = clean_code(hijo_raw)
            if canon_hijo not in self.materials_metadata:
                self.materials_metadata[canon_hijo] = MaterialMetadata(
                    canonical_hijo=canon_hijo,
                    n_parte=formatted_hijo,
                    codigo_tasa_tbar=format_part_number(get_field(row, 'Codigo TASA+TBAR', 'Codigo TASA', default=hijo_raw)),
                    descripcion=str(get_field(row, 'Descripción', 'Descripcion', 'Part Name', 'Texto breve-objeto', default='')),
                    n_interno=str(get_field(row, 'Nº Interno', 'N° Interno', 'Interno', default='')),
                    tipo_caja=str(get_field(row, 'Tipo Caja', default='')),
                    cant_x_caja=str(get_field(row, 'Cant x Caja', 'Cant. x Caja', default='1')),
                    dock=str(get_field(row, 'Dock', 'Dock/Routing', 'Routing', default='')),
                    destino=str(get_field(row, 'Destino', default='PROD')),
                    centro=str(get_field(row, 'Centro', default='1000'))
                )

    def load_materials_catalog(self, filepath: str) -> None:
        if not os.path.exists(filepath):
            return
        _, records, _ = read_table_rows(filepath, preferred_sheet_hint='material')
        for row in records:
            raw_pn = get_field(row, 'Nº Parte', 'N° Parte', 'Codigo TASA+TBAR', 'Hijo', 'Part Number', 'Material')
            canonical = clean_code(raw_pn)
            if not canonical:
                continue
            meta = self.materials_metadata.get(canonical, MaterialMetadata(canonical_hijo=canonical))
            meta.n_parte = format_part_number(meta.n_parte or raw_pn)
            meta.codigo_tasa_tbar = format_part_number(meta.codigo_tasa_tbar or get_field(row, 'Codigo TASA+TBAR', default=raw_pn))
            meta.descripcion = meta.descripcion or str(get_field(row, 'Descripción', 'Descripcion', 'Part Name', default=''))
            meta.n_interno = meta.n_interno or str(get_field(row, 'Nº Interno', 'N° Interno', default=''))
            meta.tipo_caja = meta.tipo_caja or str(get_field(row, 'Tipo Caja', default=''))
            meta.cant_x_caja = meta.cant_x_caja or str(get_field(row, 'Cant x Caja', default='1'))
            meta.dock = meta.dock or str(get_field(row, 'Dock', default=''))
            meta.destino = meta.destino or str(get_field(row, 'Destino', default='PROD'))
            meta.centro = str(get_field(row, 'Centro', default=meta.centro or '1000'))
            self.materials_metadata[canonical] = meta

    def load_format_1_file(self, filepath: str) -> None:
        if not os.path.exists(filepath):
            return
        headers, records, year = read_table_rows(filepath)
        if year:
            self.plan_year = year

        daily_cols = [h for h in headers if is_daily_column(h)]
        for col in daily_cols:
            if col not in self.daily_columns:
                self.daily_columns.append(col)

        for col in daily_cols:
            m = re.match(r'^\d{1,2}-([a-z]+)', col.strip().lower())
            if m:
                detected_m = MONTHS_LOOKUP.get(m.group(1))
                if detected_m:
                    self.firm_month = detected_m
                    break
        firm_m = getattr(self, 'firm_month', 9) or 9

        tent_cols = [h for h in headers if is_future_tentative_column(h, firm_m) and h not in self.daily_columns]
        for col in tent_cols:
            if col not in self.tentative_columns_f1:
                self.tentative_columns_f1.append(col)

        for row in records:
            raw_pn = get_field(row, 'Nº Parte', 'N° Parte', 'Part Number', 'Nº Parte Sep', 'N° Parte Sep')
            canonical = clean_code(raw_pn)
            if not canonical:
                continue
            item = self.get_or_create_demand(canonical)
            item.n_parte = format_part_number(item.n_parte or raw_pn)
            item.n_parte_sep = format_part_number(item.n_parte_sep or get_field(row, 'Nº Parte Sep', 'N° Parte Sep', default=''))
            item.codigo_tasa_tbar = format_part_number(item.codigo_tasa_tbar or get_field(row, 'Codigo TASA+TBAR', 'Codigo TASA', default=item.n_parte_sep))
            item.descripcion = item.descripcion or str(get_field(row, 'Descripción', 'Descripcion', 'Part Name', default=''))
            item.n_interno = item.n_interno or str(get_field(row, 'Nº Interno', 'N° Interno', default=''))
            item.tipo_caja = item.tipo_caja or str(get_field(row, 'Tipo Caja', default=''))
            item.cant_x_caja = item.cant_x_caja or str(get_field(row, 'Cant x Caja', default='1'))
            item.dock = item.dock or str(get_field(row, 'Dock', default=''))
            item.destino = item.destino or str(get_field(row, 'Destino', default='PROD'))

            for col in daily_cols:
                val = clean_numeric(row.get(col, 0))
                item.daily_values[col] = item.daily_values.get(col, 0.0) + val

            for col in tent_cols:
                val = clean_numeric(row.get(col, 0))
                item.tentative_values[col] = item.tentative_values.get(col, 0.0) + val

            self.register_demand_item(item)

    def load_format_2_file(self, filepath: str) -> None:
        if not os.path.exists(filepath):
            return
        headers, records, year = read_table_rows(filepath)
        if year:
            self.plan_year = year

        non_month_keys = {
            'SUPPLIER CODE', 'SUPPLIER NAME', 'PART NUMBER', 'PART NAME',
            'DOCK/ROUTING', 'KANBAN NO.', 'DOCK', 'ROUTING'
        }
        firm_m = getattr(self, 'firm_month', 9) or 9
        month_cols = [h for h in headers if h.strip().upper() not in non_month_keys and is_future_tentative_column(h, firm_m)]
        for col in month_cols:
            if col not in self.tentative_columns_f2:
                self.tentative_columns_f2.append(col)

        for row in records:
            raw_pn = get_field(row, 'Part Number', 'PART NUMBER', 'Nº Parte', 'N° Parte')
            canonical = clean_code(raw_pn)
            if not canonical:
                continue
            item = self.get_or_create_demand(canonical)
            item.descripcion = item.descripcion or str(get_field(row, 'Part Name', 'Descripción', default=''))
            item.dock = item.dock or str(get_field(row, 'Dock/Routing', 'Dock', default=''))

            for col in month_cols:
                val = clean_numeric(row.get(col, 0))
                item.tentative_values[col] = item.tentative_values.get(col, 0.0) + val

            self.register_demand_item(item)

    def get_sorted_columns(self) -> Tuple[List[str], List[str]]:
        def sort_day_key(col: str) -> int:
            m = re.match(r'^(\d{1,2})', col)
            return int(m.group(1)) if m else 999

        sorted_daily = sorted(self.daily_columns, key=sort_day_key)
        all_tentative = self.tentative_columns_f1 + [c for c in self.tentative_columns_f2 if c not in self.tentative_columns_f1]
        return sorted_daily, all_tentative

    def find_parent_demand(self, b_item: BomItem) -> Optional[ProductionItem]:
        for candidate in (
            clean_code(b_item.codigo_tasa),
            clean_code(b_item.codigo_sap_padre),
            clean_code(b_item.clave)
        ):
            if candidate and candidate in self.demand_lookup:
                return self.demand_lookup[candidate]

        for cand in (clean_code(b_item.codigo_tasa), clean_code(b_item.codigo_sap_padre)):
            if len(cand) >= 10:
                p10 = cand[:10]
                for k, it in self.demand_lookup.items():
                    if len(k) >= 10 and k[:10] == p10:
                        return it
        return None

    def get_tentative_months(self) -> List[Tuple[int, int, str, str]]:
        sorted_daily, all_tentative = self.get_sorted_columns()
        sap_tent = [col for col in all_tentative if 'TOTAL' not in col.upper()]

        firm_m = getattr(self, 'firm_month', None)
        if not firm_m:
            firm_months = set()
            for d in sorted_daily:
                m = re.match(r'^\d{1,2}-([a-z]+)', d.strip().lower())
                if m:
                    firm_months.add(MONTHS_LOOKUP.get(m.group(1), 9))
            firm_m = list(firm_months)[0] if firm_months else 9
        firm_month = firm_m

        month_names_es = {
            1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril', 5: 'Mayo', 6: 'Junio',
            7: 'Julio', 8: 'Agosto', 9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'
        }

        tentative_list: List[Tuple[int, int, str, str]] = []
        seen_m: Set[int] = set()
        for col in sap_tent:
            c_clean = col.strip().lower()
            for m_key, m_val in MONTHS_LOOKUP.items():
                if m_key in c_clean:
                    if m_val != firm_month and m_val not in seen_m:
                        seen_m.add(m_val)
                        year = self.plan_year + (1 if m_val < firm_month else 0)
                        default_date = f'01.{m_val:02d}.{year}'
                        label = f'{month_names_es.get(m_val, f"Mes {m_val}")} {year}'
                        tentative_list.append((m_val, year, default_date, label))
                    break

        return tentative_list

    def resolve_date_str(self, col_name: str) -> str:
        c_clean = col_name.strip().lower()

        m_daily = re.match(r'^(\d{1,2})-([a-z]+)', c_clean)
        if m_daily:
            day = int(m_daily.group(1))
            m_str = m_daily.group(2)
            month = MONTHS_LOOKUP.get(m_str, 9)
            year = self.plan_year
            if month < 8 and 'ago' in ''.join(self.daily_columns).lower():
                year += 1
            return f'{day:02d}.{month:02d}.{year}'

        m_num = re.match(r'^(\d{1,2})$', c_clean)
        if m_num:
            day = int(m_num.group(1))
            return f'{day:02d}.09.{self.plan_year}'

        for m_key, m_val in MONTHS_LOOKUP.items():
            if m_key in c_clean:
                if m_val in self.custom_tentative_dates:
                    return self.custom_tentative_dates[m_val]
                year = self.plan_year
                if m_val < 8:
                    year += 1
                return f'01.{m_val:02d}.{year}'

        return f'01.09.{self.plan_year}'

    def consolidate_demand(self) -> Tuple[List[str], List[List[Union[str, float]]]]:
        sorted_daily, all_tentative = self.get_sorted_columns()
        headers = [
            'Nº Parte', 'Nº Parte Sep', 'Codigo TASA+TBAR', 'PN TBAR', 'SD TBAR',
            'Descripción', 'Nº Interno', 'Tipo Caja', 'Cant x Caja', 'Dock', 'Destino'
        ] + sorted_daily + all_tentative + ['Total']

        rows: List[List[Union[str, float]]] = []
        for _, item in self.demand_items.items():
            total_sum = sum(item.daily_values.values()) + sum(item.tentative_values.values())
            row: List[Union[str, float]] = [
                item.n_parte or item.canonical_pn,
                item.n_parte_sep or item.codigo_tasa_tbar or item.n_parte,
                item.codigo_tasa_tbar or item.n_parte_sep,
                item.pn_tbar or item.canonical_pn,
                item.sd_tbar,
                item.descripcion,
                item.n_interno,
                item.tipo_caja,
                item.cant_x_caja,
                item.dock,
                item.destino,
            ]
            for col in sorted_daily:
                row.append(item.daily_values.get(col, 0.0))
            for col in all_tentative:
                row.append(item.tentative_values.get(col, 0.0))
            row.append(total_sum)
            rows.append(row)

        return headers, rows

    def consolidate_bom(self) -> Tuple[List[str], List[List[Union[str, float]]]]:
        sorted_daily, all_tentative = self.get_sorted_columns()
        all_val_cols = sorted_daily + all_tentative
        headers = ['Sector', 'Sd', 'Codigo TASA', 'Codigo SAP (Padre)', 'Hijo', 'Uso', 'Clave'] + all_val_cols + ['Total']

        rows: List[List[Union[str, float]]] = []
        for b in self.bom_items:
            row: List[Union[str, float]] = [
                b.sector,
                b.sd,
                b.codigo_tasa,
                b.codigo_sap_padre,
                b.hijo,
                b.uso,
                b.clave
            ]
            row_total = 0.0
            parent_demand = self.find_parent_demand(b)

            b.daily_calculated = {}
            for col in all_val_cols:
                if parent_demand:
                    if col in sorted_daily:
                        val = parent_demand.daily_values.get(col, 0.0)
                    else:
                        val = parent_demand.tentative_values.get(col, 0.0)
                    calculated = val * b.uso
                else:
                    calculated = 0.0
                b.daily_calculated[col] = calculated
                row.append(calculated)
                row_total += calculated

            row.append(row_total)
            rows.append(row)
        return headers, rows

    def calculate_consumption(self) -> Tuple[List[str], List[List[Union[str, float]]]]:
        self.consolidate_bom()

        sorted_daily, all_tentative = self.get_sorted_columns()
        all_val_cols = sorted_daily + all_tentative

        headers = [
            'Nº Parte', 'Codigo TASA+TBAR', 'Descripción', 'Nº Interno',
            'Tipo Caja', 'Cant x Caja', 'Dock', 'Destino'
        ] + all_val_cols + ['Total']

        unique_hijos: List[str] = []
        hijo_seen: Set[str] = set()
        for b in self.bom_items:
            canon = clean_code(b.hijo)
            if canon and canon not in hijo_seen:
                hijo_seen.add(canon)
                unique_hijos.append(b.hijo)

        rows: List[List[Union[str, float]]] = []
        for raw_hijo in unique_hijos:
            canon_hijo = clean_code(raw_hijo)
            meta = self.materials_metadata.get(canon_hijo, MaterialMetadata(canonical_hijo=canon_hijo, n_parte=raw_hijo))
            formatted_hijo = format_part_number(meta.n_parte or raw_hijo)
            row: List[Union[str, float]] = [
                formatted_hijo,
                meta.codigo_tasa_tbar or formatted_hijo,
                meta.descripcion,
                meta.n_interno,
                meta.tipo_caja,
                meta.cant_x_caja,
                meta.dock,
                meta.destino,
            ]

            hijo_total = 0.0
            matching_bom_rows = [b for b in self.bom_items if clean_code(b.hijo) == canon_hijo]

            for col in all_val_cols:
                col_sum = sum(b.daily_calculated.get(col, 0.0) for b in matching_bom_rows)
                row.append(col_sum)
                hijo_total += col_sum

            row.append(hijo_total)
            rows.append(row)

        return headers, rows

    def export_excel(self, output_filepath: str) -> None:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        themes = {
            'Demanda': {
                'title': '1. Demanda',
                'tab_color': '1F4E79',
                'header_bg': '1F4E79',
                'freeze': 'L2'
            },
            'BOM': {
                'title': '2. BOM',
                'tab_color': '2E75B6',
                'header_bg': '2E75B6',
                'freeze': 'H2'
            },
            'Consumo': {
                'title': '3. Consumo Materiales',
                'tab_color': 'C65911',
                'header_bg': 'C65911',
                'freeze': 'I2'
            }
        }

        datasets = [
            ('Demanda', self.consolidate_demand()),
            ('BOM', self.consolidate_bom()),
            ('Consumo', self.calculate_consumption())
        ]

        body_font = Font(name='Segoe UI', size=9)
        num_align = Alignment(horizontal='right', vertical='center')
        text_align = Alignment(horizontal='left', vertical='center')

        for key, (headers, rows) in datasets:
            if not headers:
                continue
            cfg = themes[key]
            ws = wb.create_sheet(title=cfg['title'])
            ws.sheet_properties.tabColor = cfg['tab_color']

            header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color=cfg['header_bg'], end_color=cfg['header_bg'], fill_type='solid')
            header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

            ws.append(headers)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align

            for r_data in rows:
                ws.append(r_data)

            for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=1, max_col=len(headers)):
                for cell in row:
                    cell.font = body_font
                    if isinstance(cell.value, (int, float)):
                        cell.alignment = num_align
                        cell.number_format = '#,##0'
                    else:
                        cell.alignment = text_align

            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    max_len = max(max_len, len(val_str))
                ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

            ws.freeze_panes = cfg['freeze']

        wb.save(output_filepath)
        wb.close()

    def export_sap_txt(self, output_txt_path: str, default_centro: str = '1000', default_version: str = '00') -> None:
        self.consolidate_bom()

        sorted_daily, all_tentative = self.get_sorted_columns()
        sap_tentative_cols = [c for c in all_tentative if 'TOTAL' not in c.upper()]
        all_val_cols = sorted_daily + sap_tentative_cols

        unique_hijos: List[str] = []
        hijo_seen: Set[str] = set()
        for b in self.bom_items:
            canon = clean_code(b.hijo)
            if canon and canon not in hijo_seen:
                hijo_seen.add(canon)
                unique_hijos.append(b.hijo)

        headers = ['Radio button Material', 'Material', 'Centro', 'Versión de plan', 'fecha', 'Tipo formato fecha', 'Cantidad']
        output_rows: List[List[str]] = []

        for raw_hijo in unique_hijos:
            canon_hijo = clean_code(raw_hijo)
            meta = self.materials_metadata.get(canon_hijo, MaterialMetadata(canonical_hijo=canon_hijo, n_parte=raw_hijo))
            mat_code = format_part_number(meta.n_parte or raw_hijo)
            centro = meta.centro or default_centro
            matching_bom_rows = [b for b in self.bom_items if clean_code(b.hijo) == canon_hijo]

            for col in all_val_cols:
                col_sum = sum(b.daily_calculated.get(col, 0.0) for b in matching_bom_rows)
                if col_sum > 0:
                    fecha_str = self.resolve_date_str(col)
                    cant_str = str(int(round(col_sum)))
                    output_rows.append([
                        'X',
                        mat_code,
                        centro,
                        default_version,
                        fecha_str,
                        'D',
                        cant_str
                    ])

        with open(output_txt_path, mode='w', encoding='utf-8', newline='') as f:
            writer = csv.writer(f, delimiter='\t', lineterminator='\r\n')
            writer.writerow(headers)
            writer.writerows(output_rows)


def find_default_file(candidates: List[str]) -> Optional[str]:
    for path in candidates:
        if os.path.exists(path):
            return path
    return None


def prompt_user_for_files() -> Tuple[List[str], List[str], List[str], List[str], str]:
    root = tk.Tk()
    root.withdraw()
    root.attributes('-topmost', True)

    excel_filetypes = [
        ('Archivos de Excel / Macros / CSV / TXT', '*.xlsx *.xlsm *.xls *.csv *.txt'),
        ('Libro de Excel habilitado para macros (*.xlsm)', '*.xlsm'),
        ('Libro de Excel (*.xlsx)', '*.xlsx'),
        ('Todos los archivos (*.*)', '*.*')
    ]

    auto_bom = find_default_file(DEFAULT_BOM_CANDIDATES)
    if auto_bom:
        bom_files = [auto_bom]
    else:
        messagebox.showinfo(
            'Consolidador de Producción',
            'Paso 1: Selecciona el archivo de BOM / Recetas (.xlsm, .xlsx, .csv o .txt).'
        )
        bom_file = filedialog.askopenfilename(
            title='Selecciona el archivo de BOM (Recetas)',
            filetypes=excel_filetypes
        )
        bom_files = [bom_file] if bom_file else []

    auto_materials = find_default_file(DEFAULT_MATERIALS_CANDIDATES)
    if auto_materials:
        mapping_files = [auto_materials]
    else:
        messagebox.showinfo(
            'Consolidador de Producción',
            'Paso 2: Selecciona el archivo de Mapeo / Catálogo de Materiales (opcional).'
        )
        mapping_file = filedialog.askopenfilename(
            title='Selecciona el archivo de Mapeo / Materiales (opcional)',
            filetypes=excel_filetypes
        )
        mapping_files = [mapping_file] if mapping_file else []

    messagebox.showinfo(
        'Consolidador de Producción',
        'Selecciona los archivos de Formato 1 (Demanda Diaria - selección múltiple).'
    )
    format_1_files = list(filedialog.askopenfilenames(
        title='Selecciona los archivos de Formato 1 (Diario)',
        filetypes=excel_filetypes
    ))

    messagebox.showinfo(
        'Consolidador de Producción',
        'Selecciona los archivos de Formato 2 (Mediano Plazo / Proveedores - selección múltiple).'
    )
    format_2_files = list(filedialog.askopenfilenames(
        title='Selecciona los archivos de Formato 2 (Mediano Plazo)',
        filetypes=excel_filetypes
    ))

    if not format_1_files and not format_2_files:
        messagebox.showwarning('Cancelado', 'No se seleccionaron archivos de producción.')
        return [], [], [], [], ''

    output_path = filedialog.asksaveasfilename(
        title='Guardar Libro Consolidado como...',
        defaultextension='.xlsx',
        initialfile='Consolidado_Demanda_BOM_Consumo.xlsx',
        filetypes=[('Libro de Excel (*.xlsx)', '*.xlsx')]
    )

    return bom_files, mapping_files, format_1_files, format_2_files, output_path


def prompt_tentative_dates_dialog(tentative_info: List[Tuple[int, int, str, str]]) -> Dict[int, str]:
    if not tentative_info:
        return {}

    dialog = tk.Toplevel()
    dialog.title('Fechas de Carga para Meses Tentativos')
    dialog.attributes('-topmost', True)
    dialog.resizable(False, False)
    dialog.grab_set()

    title_label = tk.Label(
        dialog,
        text='Seleccione o ingrese la fecha para cada mes no firme (DD.MM.YYYY):',
        font=('Segoe UI', 10, 'bold')
    )
    title_label.pack(padx=20, pady=(15, 10))

    frame = tk.Frame(dialog)
    frame.pack(padx=20, pady=5, fill='x')

    entries: Dict[int, tk.Entry] = {}
    for m_val, year, default_date, label in tentative_info:
        row_frame = tk.Frame(frame)
        row_frame.pack(fill='x', pady=4)
        lbl = tk.Label(row_frame, text=f'{label}:', width=22, anchor='w', font=('Segoe UI', 9))
        lbl.pack(side='left')
        ent = tk.Entry(row_frame, width=16, font=('Segoe UI', 9))
        ent.insert(0, default_date)
        ent.pack(side='right')
        entries[m_val] = ent

    result: Dict[int, str] = {}

    def on_confirm() -> None:
        for m_val, ent in entries.items():
            val = ent.get().strip().replace('/', '.').replace('-', '.')
            result[m_val] = val
        dialog.destroy()

    btn = tk.Button(
        dialog,
        text='Confirmar y Generar Archivos',
        command=on_confirm,
        bg='#1F4E79',
        fg='white',
        font=('Segoe UI', 9, 'bold'),
        padx=12,
        pady=5,
        relief='flat'
    )
    btn.pack(pady=(15, 18))

    dialog.wait_window()
    return result


def parse_tentative_dates_arg(raw_args: List[str], tentative_months_info: List[Tuple[int, int, str, str]]) -> Dict[int, str]:
    res: Dict[int, str] = {}
    if not raw_args:
        return res

    for item in raw_args:
        if ',' in item and '=' not in item:
            parts = [p.strip() for p in item.split(',') if p.strip()]
            for idx, p in enumerate(parts):
                if idx < len(tentative_months_info):
                    m_val = tentative_months_info[idx][0]
                    res[m_val] = p.replace('/', '.').replace('-', '.')
        elif '=' in item:
            k, v = item.split('=', 1)
            k = k.strip().lower()
            v = v.strip().replace('/', '.').replace('-', '.')
            if k.isdigit():
                res[int(k)] = v
            elif k in MONTHS_LOOKUP:
                res[MONTHS_LOOKUP[k]] = v
        else:
            v = item.strip().replace('/', '.').replace('-', '.')
            if len(res) < len(tentative_months_info):
                m_val = tentative_months_info[len(res)][0]
                res[m_val] = v
    return res


def main() -> None:
    parser = argparse.ArgumentParser(description='Consolidador de Demanda, BOM y Consumo de Materiales')
    parser.add_argument('--bom', nargs='*', default=[], help='Ruta archivo de BOM (.xlsm / .xlsx)')
    parser.add_argument('--mapping', nargs='*', default=[], help='Ruta archivo de mapeo')
    parser.add_argument('--f1', nargs='*', default=[], help='Rutas archivos Formato 1')
    parser.add_argument('--f2', nargs='*', default=[], help='Rutas archivos Formato 2')
    parser.add_argument('--tentative-dates', nargs='*', default=[], help='Fechas personalizadas para meses no firmes (ej: 10=01.10.2026 o lista)')
    parser.add_argument('--output', default='', help='Ruta archivo de salida .xlsx')
    parser.add_argument('--output-txt', default='', help='Ruta archivo de salida TXT SAP')

    args = parser.parse_args()

    if args.f1 or args.f2:
        bom_files = args.bom or ([find_default_file(DEFAULT_BOM_CANDIDATES)] if find_default_file(DEFAULT_BOM_CANDIDATES) else [])
        mapping_files = args.mapping or ([find_default_file(DEFAULT_MATERIALS_CANDIDATES)] if find_default_file(DEFAULT_MATERIALS_CANDIDATES) else [])
        format_1_files = args.f1
        format_2_files = args.f2
        output_path = args.output or 'Consolidado_Demanda_BOM_Consumo.xlsx'
    else:
        bom_files, mapping_files, format_1_files, format_2_files, output_path = prompt_user_for_files()

    if not output_path:
        return

    consolidator = ProductionConsolidator()
    for mf in mapping_files:
        consolidator.load_mapping_file(mf)
        consolidator.load_materials_catalog(mf)
    for bf in bom_files:
        consolidator.load_bom_file(bf)
    for f1 in format_1_files:
        consolidator.load_format_1_file(f1)
    for f2 in format_2_files:
        consolidator.load_format_2_file(f2)

    tentative_info = consolidator.get_tentative_months()
    if args.tentative_dates:
        custom_dates = parse_tentative_dates_arg(args.tentative_dates, tentative_info)
        consolidator.custom_tentative_dates.update(custom_dates)
    elif not (args.f1 or args.f2) and tk._default_root:
        custom_dates = prompt_tentative_dates_dialog(tentative_info)
        consolidator.custom_tentative_dates.update(custom_dates)

    consolidator.export_excel(output_path)

    base_dir = os.path.dirname(output_path)
    base_name = os.path.splitext(os.path.basename(output_path))[0]
    txt_path = args.output_txt or os.path.join(base_dir, f'{base_name}_SAP_PIR.txt')
    consolidator.export_sap_txt(txt_path)

    if tk._default_root:
        messagebox.showinfo(
            'Completado',
            f'Archivos generados exitosamente:\n\n1. Libro Excel (3 pestañas):\n{output_path}\n\n2. Carga SAP (TXT delimitado por tabulaciones):\n{txt_path}'
        )


if __name__ == '__main__':
    main()
