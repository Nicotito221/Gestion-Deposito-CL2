import io
import re
from datetime import date, timedelta
from typing import Any, Dict, List, Optional, Tuple
import openpyxl

DAY_KEYWORDS: List[Tuple[str, int, re.Pattern]] = [
    ("LUNES", 0, re.compile(r"\b(LUNES|LUN)\b", re.IGNORECASE)),
    ("MARTES", 1, re.compile(r"\b(MARTES|MAR)\b", re.IGNORECASE)),
    ("MIERCOLES", 2, re.compile(r"\b(MIERCOLES|MI[EÉ]RCOLES|MIE)\b", re.IGNORECASE)),
    ("JUEVES", 3, re.compile(r"\b(JUEVES|JUE)\b", re.IGNORECASE)),
    ("VIERNES", 4, re.compile(r"\b(VIERNES|VIE)\b", re.IGNORECASE)),
    ("SABADO", 5, re.compile(r"\b(SABADO|S[AÁ]BADO|SAB)\b", re.IGNORECASE)),
    ("DOMINGO", 6, re.compile(r"\b(DOMINGO|DOM)\b", re.IGNORECASE)),
]

RANGE_REGEX = re.compile(
    r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})\s*(?:al|a|-)\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})",
    re.IGNORECASE,
)

SINGLE_DATE_REGEX = re.compile(
    r"(?:SEMANA|DEL|FECHA|DESDE)\s*[:.]*\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})",
    re.IGNORECASE,
)

CELL_DATE_REGEX = re.compile(r"(\d{1,2})[/.-](\d{1,2})(?:[/.-](\d{2,4}))?")


def detect_day_in_text(text: str) -> Optional[Tuple[str, int, Optional[Tuple[int, int, Optional[int]]]]]:
    clean = text.strip()
    for norm_name, day_idx, pat in DAY_KEYWORDS:
        if pat.search(clean):
            d_match = CELL_DATE_REGEX.search(clean)
            date_tuple = None
            if d_match:
                d_day = int(d_match.group(1))
                d_month = int(d_match.group(2))
                d_year = int(d_match.group(3)) if d_match.group(3) else None
                if d_year and d_year < 100:
                    d_year += 2000
                date_tuple = (d_day, d_month, d_year)
            return norm_name, day_idx, date_tuple
    return None


def parse_plan_semanal_excel(
    file_bytes: bytes,
    override_fecha_inicio: Optional[date] = None,
) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    sheet_names = [s.lower() for s in wb.sheetnames]
    ws = None
    for target in ["resumen", "plan"]:
        if target in sheet_names:
            ws = wb[wb.sheetnames[sheet_names.index(target)]]
            break
    if ws is None:
        ws = wb.active

    fecha_inicio: Optional[date] = override_fecha_inicio
    fecha_fin: Optional[date] = None

    if fecha_inicio is None:
        for r in range(1, 10):
            for c in range(1, 30):
                val = ws.cell(row=r, column=c).value
                if val and isinstance(val, str):
                    m = RANGE_REGEX.search(val)
                    if m:
                        d1, m1, y1, d2, m2, y2 = m.groups()
                        y1_int = int(y1) + 2000 if len(y1) <= 2 else int(y1)
                        y2_int = int(y2) + 2000 if len(y2) <= 2 else int(y2)
                        fecha_inicio = date(y1_int, int(m1), int(d1))
                        fecha_fin = date(y2_int, int(m2), int(d2))
                        break
                    m_single = SINGLE_DATE_REGEX.search(val)
                    if m_single:
                        d1, m1, y1 = m_single.groups()
                        y1_int = int(y1) + 2000 if len(y1) <= 2 else int(y1)
                        fecha_inicio = date(y1_int, int(m1), int(d1))
                        fecha_fin = fecha_inicio + timedelta(days=5)
                        break
            if fecha_inicio:
                break

    merged_map: Dict[Tuple[int, int], Tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        merged_map[(rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)

    header_row_idx: Optional[int] = None
    raw_days: List[Dict[str, Any]] = []

    for r in range(2, 10):
        found = []
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val and isinstance(val, str):
                info = detect_day_in_text(val)
                if info:
                    norm_name, day_idx, date_tuple = info
                    if date_tuple is None and r + 1 <= ws.max_row:
                        next_val = ws.cell(row=r + 1, column=c).value
                        if next_val and isinstance(next_val, str):
                            d_m = CELL_DATE_REGEX.search(next_val)
                            if d_m:
                                d_d = int(d_m.group(1))
                                d_mo = int(d_m.group(2))
                                d_y = int(d_m.group(3)) if d_m.group(3) else None
                                if d_y and d_y < 100:
                                    d_y += 2000
                                date_tuple = (d_d, d_mo, d_y)

                    max_c = c
                    if (r, c) in merged_map:
                        max_c = merged_map[(r, c)][1]
                    found.append({
                        "nombre": norm_name,
                        "day_idx": day_idx,
                        "date_tuple": date_tuple,
                        "min_col": c,
                        "max_col": max_c,
                    })
        if len(found) >= 3:
            header_row_idx = r
            raw_days = found
            break

    if not raw_days or header_row_idx is None:
        raise ValueError("No se encontraron las columnas de los días en la estructura del Excel.")

    if fecha_inicio is None:
        first_d = raw_days[0]
        if first_d.get("date_tuple"):
            d_d, d_m, d_y = first_d["date_tuple"]
            y_use = d_y if d_y else 2026
            fecha_inicio = date(y_use, d_m, d_d)
        else:
            raise ValueError("No se pudo detectar el rango de fechas en la cabecera del archivo.")

    if fecha_fin is None:
        fecha_fin = fecha_inicio + timedelta(days=5)

    raw_days.sort(key=lambda x: x["min_col"])

    day_columns: List[Dict[str, Any]] = []
    for i, rd in enumerate(raw_days):
        if rd.get("date_tuple"):
            d_day, d_month, d_year = rd["date_tuple"]
            if not d_year:
                if d_month < fecha_inicio.month and fecha_inicio.month == 12:
                    d_year = fecha_inicio.year + 1
                else:
                    d_year = fecha_inicio.year
            col_date = date(d_year, d_month, d_day)
        else:
            col_date = fecha_inicio + timedelta(days=rd["day_idx"])

        if col_date < fecha_inicio or col_date > fecha_fin:
            continue
        if any(d["fecha"] == col_date for d in day_columns):
            continue

        min_c = rd["min_col"]
        if i + 1 < len(raw_days):
            next_min_c = raw_days[i + 1]["min_col"]
            max_c = max(min_c, next_min_c - 1)
        else:
            max_c = max(rd["max_col"], min_c + 2)

        col_span = list(range(min_c, max_c + 1))

        day_columns.append({
            "nombre": rd["nombre"],
            "fecha": col_date,
            "columns": col_span,
        })

    if not day_columns:
        raise ValueError("No se pudieron validar los días dentro del rango semanal especificado.")

    sd_col_idx: Optional[int] = None
    for c in range(1, 6):
        for r in range(max(1, header_row_idx - 2), header_row_idx + 2):
            val = ws.cell(row=r, column=c).value
            if val and isinstance(val, str) and any(k in val.upper() for k in ["PROD", "SD", "PART", "CODIGO"]):
                sd_col_idx = c
                break
        if sd_col_idx:
            break
    if sd_col_idx is None:
        sd_col_idx = 3

    filas: List[Dict[str, Any]] = []
    totales_dia = {d["fecha"].isoformat(): 0.0 for d in day_columns}

    start_row = header_row_idx + 1
    shift_val = ws.cell(row=start_row, column=day_columns[0]["columns"][0]).value
    if shift_val and isinstance(shift_val, str) and any(k in shift_val.upper() for k in ["SHIFT", "TURNO", "1ST", "T1"]):
        start_row += 1

    stop_keywords = ["TOTAL", "DETALLE", "CAMBIO", "HORA", "OBSERV"]

    for r in range(start_row, ws.max_row + 1):
        sd_raw = ws.cell(row=r, column=sd_col_idx).value
        if sd_raw is None:
            continue
        sd_str = str(sd_raw).strip().upper()
        if not sd_str or any(k in sd_str for k in stop_keywords):
            if "TOTAL" in sd_str:
                break
            continue

        cantidades: Dict[str, float] = {}
        row_total = 0.0
        has_any_num = False
        for day in day_columns:
            day_sum = 0.0
            for col in day["columns"]:
                c_val = ws.cell(row=r, column=col).value
                if isinstance(c_val, (int, float)):
                    day_sum += float(c_val)
                    has_any_num = True
                elif isinstance(c_val, str):
                    try:
                        clean_num = float(c_val.replace(",", ".").strip())
                        day_sum += clean_num
                        has_any_num = True
                    except ValueError:
                        pass
            iso_d = day["fecha"].isoformat()
            cantidades[iso_d] = round(day_sum, 2)
            row_total += day_sum
            totales_dia[iso_d] += day_sum

        if has_any_num or sd_str.startswith("H"):
            filas.append({
                "sd": sd_str,
                "cantidades": cantidades,
                "total_semana": round(row_total, 2),
            })

    total_general = sum(totales_dia.values())

    return {
        "fecha_inicio": fecha_inicio.isoformat(),
        "fecha_fin": fecha_fin.isoformat(),
        "dias": [
            {"nombre": d["nombre"], "fecha": d["fecha"].isoformat(), "total": round(totales_dia[d["fecha"].isoformat()], 2)}
            for d in day_columns
        ],
        "filas": filas,
        "totales_dia": {k: round(v, 2) for k, v in totales_dia.items()},
        "total_general": round(total_general, 2),
    }
