import csv
import io
import os
from decimal import Decimal, InvalidOperation
from typing import BinaryIO, Dict, Iterator, List, Optional, Tuple, Union

import openpyxl

from app.schemas.bom import BomParseResult, BomRelationRecord, MaterialRecord


def clean_quantity(raw: Union[str, float, int, None]) -> Decimal:
    if raw is None:
        return Decimal("1.0000")
    if isinstance(raw, (int, float)):
        return Decimal(str(raw))

    s = str(raw).strip()
    if not s or s in ("-", "N/A", "NULL", "NONE"):
        return Decimal("1.0000")

    if "." in s and "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")

    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal("1.0000")


def is_header_row(cols: List[str]) -> bool:
    for c in cols[:6]:
        cleaned = c.strip().upper()
        if cleaned in ("MATERIAL", "PARTE", "CODIGO", "Nº PARTE", "N° PARTE", "PART NUMBER", "HIJO"):
            return True
    return False


def extract_row_data(cols: List[str]) -> Optional[Tuple[int, str, str, Decimal, str]]:
    part_number: str = ""
    level: int = -1

    for idx in range(min(6, len(cols))):
        val = cols[idx].strip()
        if val:
            part_number = val.upper()
            level = idx
            break

    if not part_number or level == -1:
        return None

    if len(part_number) > 30:
        part_number = part_number[:30]

    desc: str = ""
    if len(cols) > 6 and cols[6].strip():
        desc = cols[6].strip()[:100]

    qty: Decimal = Decimal("1.0000")
    if len(cols) > 7 and cols[7].strip():
        qty = clean_quantity(cols[7])

    unit: str = "UN"
    if len(cols) > 8 and cols[8].strip():
        unit = cols[8].strip()[:10]

    return level, part_number, desc, qty, unit


class BomParser:
    @staticmethod
    def parse(source: Union[str, BinaryIO, io.BytesIO]) -> BomParseResult:
        materiales: Dict[str, MaterialRecord] = {}
        relaciones: List[BomRelationRecord] = []
        current_stack: Dict[int, str] = {}
        current_root_product: Optional[str] = None
        total_lineas = 0
        orden_secuencia = 0

        rows_iterator = BomParser._get_row_iterator(source)

        for raw_cols in rows_iterator:
            total_lineas += 1
            if not raw_cols:
                continue

            str_cols = [str(c) if c is not None else "" for c in raw_cols]
            if is_header_row(str_cols):
                continue

            extracted = extract_row_data(str_cols)
            if extracted is None:
                continue

            level, part_number, desc, qty, unit = extracted

            if part_number not in materiales:
                materiales[part_number] = MaterialRecord(
                    id_material=part_number,
                    nombre=desc or part_number,
                    unidad=unit or "UN"
                )
            elif desc and materiales[part_number].nombre == part_number:
                materiales[part_number].nombre = desc

            if level == 0:
                current_root_product = part_number
                current_stack = {0: part_number}
            else:
                if not current_root_product:
                    current_root_product = current_stack.get(0, part_number)

                id_padre = current_stack.get(level - 1)
                if not id_padre:
                    for l in range(level - 1, -1, -1):
                        if l in current_stack:
                            id_padre = current_stack[l]
                            break

                if id_padre and id_padre != part_number:
                    orden_secuencia += 1
                    relaciones.append(
                        BomRelationRecord(
                            id_producto_raiz=current_root_product,
                            id_padre=id_padre,
                            id_hijo=part_number,
                            cantidad=qty,
                            unidad=unit,
                            nivel=level,
                            orden=orden_secuencia
                        )
                    )

                current_stack[level] = part_number
                current_stack = {k: v for k, v in current_stack.items() if k <= level}

        return BomParseResult(
            materiales=materiales,
            relaciones=relaciones,
            total_lineas_archivo=total_lineas
        )

    @staticmethod
    def _get_row_iterator(source: Union[str, BinaryIO, io.BytesIO]) -> Iterator[List[str]]:
        if isinstance(source, str):
            with open(source, "rb") as f:
                header_bytes = f.read(4)
                f.seek(0)
                if header_bytes.startswith(b"PK\x03\x04"):
                    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
                    ws = wb.active
                    for row in ws.iter_rows(values_only=True):
                        yield [str(c) if c is not None else "" for c in row]
                    wb.close()
                    return

            encodings = ["utf-8", "latin-1", "cp1252"]
            for enc in encodings:
                try:
                    with open(source, mode="r", encoding=enc) as f:
                        reader = csv.reader(f, delimiter="\t")
                        for row in reader:
                            yield row
                    return
                except UnicodeDecodeError:
                    continue
            raise ValueError(f"No se pudo decodificar el archivo {source} con codificaciones estándar.")

        else:
            header_bytes = source.read(4)
            source.seek(0)
            if header_bytes.startswith(b"PK\x03\x04"):
                wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
                ws = wb.active
                for row in ws.iter_rows(values_only=True):
                    yield [str(c) if c is not None else "" for c in row]
                wb.close()
                return

            content = source.read()
            text_data = None
            for enc in ["utf-8", "latin-1", "cp1252"]:
                try:
                    text_data = content.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue

            if text_data is None:
                raise ValueError("No se pudo decodificar el stream binario.")

            reader = csv.reader(io.StringIO(text_data), delimiter="\t")
            for row in reader:
                yield row
